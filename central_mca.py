import csv
import io
import math
import random
import statistics
import sys
import threading
import time
from collections import deque
from datetime import datetime

import serial
import serial.tools.list_ports
from dash import Dash, html, dcc, Input, Output, State, ALL, ctx, no_update
from dash.exceptions import PreventUpdate
import plotly.graph_objects as go
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------
# COMO EXECUTAR
#   python central_mca.py            -> modo normal (lê a placa pela serial)
#   python central_mca.py --teste    -> placa SIMULADA (para testar sem sensores)
#   python central_mca.py --diag     -> modo normal + mostra no terminal cada linha recebida
#   python central_mca.py --calibrar -> mede o R0 do MQ-135 em ar limpo (5 min) e imprime o valor
# ---------------------------------------------------------------
MODO_TESTE = "--teste" in sys.argv
DIAGNOSTICO_SERIAL = "--diag" in sys.argv
CALIBRAR = "--calibrar" in sys.argv

# ---------------------------------------------------------------
# CONFIGURAÇÃO
# ---------------------------------------------------------------
MAX_PONTOS = 120        # quantos pontos o gráfico guarda
MAX_OCORRENCIAS = 1000  # quantas ocorrências o histórico guarda (antes: 50)
MAX_TRILHA = 100        # quantas posições recentes do GPS guardar no trajeto
VALIDADE_GPS_S = 10     # segundos até considerar os dados do GPS desatualizados
INTERVALO_S = 1         # segundos entre leituras simuladas
BAUD = 9600             # precisa ser igual ao Serial.begin() do sketch da Mega
PORTA_MANUAL = None     # ex.: "COM5". Use para o HC-05 ou se a placa não for CH340.

# ---------------------------------------------------------------
# MQ-135 (GÁS): do número do ADC até o ppm estimado
#
# Passo a passo:
#   1) ADC (0 a 1023) -> resistência do sensor Rs:   Rs = RL * (1023 - ADC) / ADC
#   2) Rs / R0 = "razão". R0 é o Rs do SEU sensor em ar limpo dividido por 3,6.
#   3) ppm = a * razão^b   (curva ajustada sobre o gráfico do datasheet; cada gás tem a e b)
#
# IMPORTANTE (honestidade técnica):
#   - O MQ-135 tem sensibilidade cruzada: ele NÃO sabe qual gás é. O mesmo Rs gera um ppm
#     "equivalente" diferente para cada gás. O que a central mostra é uma ESTIMATIVA.
#   - Sem calibrar o R0 em ar limpo, o ppm não faz sentido. Rode:  python central_mca.py --calibrar
#   - Temperatura e umidade alteram a leitura (esta versão não corrige isso).
#   - Protótipo escolar: NÃO é instrumento certificado de segurança.
# ---------------------------------------------------------------
ADC_MAXIMO = 1023
MQ135_RL_KOHM = 10.0     # resistor de carga do módulo (em kΩ). Olhe o SMD perto do sensor: "103" = 10 kΩ,
                         # "102" = 1 kΩ, "223" = 22 kΩ. Se for diferente de 10, mude aqui.
MQ135_R0_KOHM = None     # COLOQUE AQUI o valor impresso por --calibrar (ex.: 18.4). None = sem calibração.
MQ135_R0_KOHM_SIMULADO = 20.0   # R0 de mentira, usado SÓ no --teste
MQ135_RAZAO_AR_LIMPO = 3.6      # Rs/R0 em ar limpo (valor usado pela biblioteca MQSensorsLib)
MQ135_GAS_PRINCIPAL = "amonia"  # gás cujo "ppm equivalente" aparece no cartão e no gráfico
MQ135_FRACAO_ATENCAO = 0.5      # ATENÇÃO quando algum gás candidato chega a 50% do limite dele
PPM_MAXIMO_EXIBIDO = 10000.0    # fora da faixa do datasheet a curva é extrapolação: limita o número

# Cada gás candidato: coeficientes da curva (a, b), limite de tolerância (NR-15, Anexo 11) e
# "fundo" (concentração normal no ar, somada ao resultado; só o CO2 usa).
# CONFERIDOS: amônia, álcool e CO2 (coeficientes) e amônia 20 / CO 39 / CO2 3900 / acetona 780 /
#   álcool etílico 780 ppm (limites NR-15).
# CONFERIR NO TEXTO OFICIAL: coeficientes de CO, tolueno e acetona; limite do tolueno (78 ppm).
# O Anexo 11 da NR-15 está em processo de revisão: se mudar, altere "limite_ppm" aqui.
GASES_MQ135 = {
    "amonia":  {"nome": "Amônia (NH₃)", "curto": "NH₃",     "a": 102.2,  "b": -2.473, "limite_ppm": 20,   "fundo": 0},
    "alcool":  {"nome": "Álcool",       "curto": "álcool",  "a": 77.255, "b": -3.18,  "limite_ppm": 780,  "fundo": 0},
    "co":      {"nome": "CO",           "curto": "CO",      "a": 605.18, "b": -3.937, "limite_ppm": 39,   "fundo": 0},
    "co2":     {"nome": "CO₂",          "curto": "CO₂",     "a": 110.47, "b": -2.862, "limite_ppm": 3900, "fundo": 400},
    "tolueno": {"nome": "Tolueno",      "curto": "tolueno", "a": 44.947, "b": -3.445, "limite_ppm": 78,   "fundo": 0},
    "acetona": {"nome": "Acetona",      "curto": "acetona", "a": 34.668, "b": -3.369, "limite_ppm": 780,  "fundo": 0},
}


def ppm_do_gas(gas, razao):
    """ppm estimado de um gás a partir da razão Rs/R0 (curva a * razão^b)."""
    c = GASES_MQ135[gas]
    razao = max(razao, 0.1)   # evita números absurdos se a razão chegar perto de zero
    return min(c["a"] * razao ** c["b"] + c["fundo"], PPM_MAXIMO_EXIBIDO)


def razao_para_ppm(gas, ppm_alvo):
    """Operação inversa: qual razão Rs/R0 produz esse ppm para esse gás."""
    c = GASES_MQ135[gas]
    base = (ppm_alvo - c["fundo"]) / c["a"]
    if base <= 0:
        return float("inf")
    return base ** (1.0 / c["b"])


def razao_no_limite(fracao):
    """Razão Rs/R0 em que ALGUM gás candidato chega a 'fracao' do seu limite.
    Rs/R0 cai quando o gás aumenta; então o MAIOR valor é o primeiro a ser alcançado (pior caso)."""
    candidatas = [razao_para_ppm(g, fracao * c["limite_ppm"]) for g, c in GASES_MQ135.items()]
    return max(r for r in candidatas if math.isfinite(r))


# Duas "linhas" da razão: abaixo da primeira é ATENÇÃO, abaixo da segunda é ALERTA (risco).
RAZAO_ATENCAO = razao_no_limite(MQ135_FRACAO_ATENCAO)
RAZAO_ALERTA = razao_no_limite(1.0)
# As mesmas linhas, convertidas para ppm equivalente do gás principal (para o gráfico).
LIMITE_GAS_PPM = round(ppm_do_gas(MQ135_GAS_PRINCIPAL, RAZAO_ALERTA), 1)
ATENCAO_GAS_PPM = round(ppm_do_gas(MQ135_GAS_PRINCIPAL, RAZAO_ATENCAO), 1)

# Traduz o nome curto que a placa envia para o id do sensor no cadastro.
MAPA_SERIAL = {
    "mq135": "mq135",
    "lux": "luminosidade",
    "temp": "temperatura",   # DS18B20: o Arduino envia "temp:27.3"
    "bpm": "frequencia",
    "spo2": "spo2",
}

# OPCIONAL: campos de "o módulo está presente?" que a placa pode enviar.
# Se o sketch enviar "max:1" (MAX30102 respondendo) ou "max:0" (não respondendo),
# a central esconde BPM e SpO₂ na hora quando vier "max:0". Se a placa não enviar
# esse campo, nada muda: vale só o tempo sem dados (validade_s).
MODULOS_SERIAL = {
    "max": ("frequencia", "spo2"),
}

# Cadastro dos sensores. Para adicionar um novo, acrescente um dicionário aqui.
# "casas"       = quantas casas decimais mostrar no cartão.
# "origem"      = "simulado" (demonstração), "serial" (vem da Mega), "gps" (campos do GPS)
#                 ou None (ainda não ligado).
# "faixa"       = (mínimo, máximo) fixo do eixo Y do gráfico (opcional).
# "limites"     = {"min": valor ou None, "max": valor ou None} (opcional).
#                 Fora desse intervalo, o sensor entra em alerta.
# "atencao_max" = (opcional) linha de ATENÇÃO desenhada no gráfico (hoje só o gás usa).
# "eixo_zero"   = (opcional) força o eixo Y do gráfico a começar em zero.
# "validade_s"  = depois de quantos segundos sem leitura nova o sensor é dado como
#                 "sem dados" e some da Conexão (opcional).
# "ocioso"      = valor "de espera" (ex.: 0 bpm sem dedo). Se a placa envia um valor
#                 menor ou igual a ele, OU se a placa avisa que o módulo está presente (max:1)
#                 mas não manda valor, o cartão fica em "Sem leitura" e nunca gera alerta.
# "faixa_valida"= (mínimo, máximo) que o sensor consegue medir de verdade. Leituras fora
#                 dela são DESCARTADAS (ex.: o DS18B20 desconectado devolve -127 °C).
SENSORES = [
    {"id": "luminosidade", "nome": "Luminosidade", "descricao": "BH1750 (limites provisórios)",
     "unidade": "lux", "casas": 0, "origem": "serial",
     # PROVISÓRIO: troque pelos valores definidos com seu professor / norma de iluminação.
     "limites": {"min": 100, "max": 10000}, "validade_s": 5},
    {"id": "frequencia", "nome": "Frequência cardíaca",
     "descricao": "MAX30102 (estimativa, apoie o dedo) • limites provisórios",
     "unidade": "bpm", "casas": 0, "origem": "serial", "faixa": (40, 200),
     # PROVISÓRIO: faixa de repouso de um adulto. Ajuste com seu professor.
     "limites": {"min": 50, "max": 120},
     "validade_s": 5, "ocioso": {"valor": 0}},
    {"id": "spo2", "nome": "SpO₂",
     "descricao": "MAX30102 (estimativa, não calibrado) • limite provisório",
     "unidade": "%", "casas": 0, "origem": "serial", "faixa": (70, 100),
     # PROVISÓRIO: só limite mínimo.
     "limites": {"min": 92, "max": None},
     "validade_s": 5, "ocioso": {"valor": 0}},
    {"id": "mq135", "nome": "Gás",
     "descricao": "MQ-135 • ppm estimado pela curva do datasheet (equivalente de amônia), requer calibração do R0 em ar limpo",
     "unidade": "ppm", "casas": 1, "origem": "serial", "eixo_zero": True,
     # Linhas calculadas dos limites da NR-15 (pior caso entre os gases candidatos).
     "limites": {"min": None, "max": LIMITE_GAS_PPM}, "atencao_max": ATENCAO_GAS_PPM,
     "validade_s": 5},
    {"id": "temperatura", "nome": "Temperatura",
     "descricao": "DS18B20 • temperatura no ponto de contato • limites provisórios",
     "unidade": "°C", "casas": 1, "origem": "serial",
     # PROVISÓRIO: depende do que o DS18B20 mede (ambiente ou corpo). Ajuste.
     "limites": {"min": 15, "max": 38},
     "validade_s": 5, "faixa_valida": (-55, 125)},   # faixa de medição do DS18B20
    {"id": "gps", "nome": "GPS", "descricao": "NEO-6M (precisa de céu aberto)",
     "unidade": "", "casas": 0, "origem": "gps"},
]
SENSOR_POR_ID = {s["id"]: s for s in SENSORES}

# Estados possíveis de um cartão: texto exibido e cor.
ESTADOS = {
    "sem_dados": {"texto": "Sem dados", "cor": "#94a3b8"},
    "ocioso": {"texto": "Sem leitura (sem dedo ou medindo)", "cor": "#94a3b8"},
    "simulado": {"texto": "Simulado (demonstração)", "cor": "#D4AF37"},
    "real": {"texto": "Dados reais", "cor": "#D98B55"},
    "normal": {"texto": "Dentro do limite", "cor": "#D98B55"},
    "alerta_baixo": {"texto": "ALERTA: abaixo do mínimo", "cor": "#fb7185"},
    "alerta_alto": {"texto": "ALERTA: acima do máximo", "cor": "#fb7185"},
    "desconectado": {"texto": "Placa desconectada", "cor": "#fb7185"},
    "aquecendo": {"texto": "Aquecendo (não avalie ainda)", "cor": "#f59e0b"},
    "aprendendo": {"texto": "Aguardando leitura", "cor": "#f59e0b"},
    "sem_calibracao": {"texto": "Sem calibração (rode --calibrar)", "cor": "#f59e0b"},
    "ar_normal": {"texto": "Gás: abaixo do limite", "cor": "#34d399"},
    "ar_atencao": {"texto": "ATENÇÃO: gás se aproximando do limite", "cor": "#D4AF37"},
    "ar_alerta": {"texto": "ALERTA: gás acima do limite (risco)", "cor": "#fb7185"},
    "gps_ok": {"texto": "Posição válida", "cor": "#D98B55"},
    "gps_sem_sinal": {"texto": "Sem sinal de satélite", "cor": "#D4AF37"},
    "gps_sem_modulo": {"texto": "Módulo GPS sem dados", "cor": "#fb7185"},
}
ESTADOS_DE_ALERTA = ("alerta_baixo", "alerta_alto", "ar_atencao", "ar_alerta")
ESTADOS_ATENCAO = ("ar_atencao",)
ESTADOS_NORMAIS = ("normal", "ar_normal")

# Estados em que o sensor NÃO está enviando dados: o cartão some e ele sai da lista de conexão.
# ("ocioso" não entra aqui: a placa está falando, só não há leitura, como BPM sem dedo.)
ESTADOS_INATIVOS = ("desconectado", "sem_dados", "gps_sem_modulo")

# Cada estado vira uma categoria; o CSS (.est-ok, .est-alerta...) decide a cor.
CATEGORIA = {
    "normal": "ok", "ar_normal": "ok", "real": "ok", "gps_ok": "ok",
    "ar_atencao": "atencao", "aquecendo": "atencao", "aprendendo": "atencao", "gps_sem_sinal": "atencao",
    "sem_calibracao": "atencao",
    "alerta_baixo": "alerta", "alerta_alto": "alerta", "ar_alerta": "alerta",
    "simulado": "demo",
}  # o que não está aqui (sem_dados, ocioso, desconectado...) cai em "neutro" (cinza)

ESTIMADOS = {"mq135", "frequencia", "spo2"}   # ganham o selo EST. no cartão

# ---------------------------------------------------------------
# ARMAZENAMENTO DOS DADOS
# ---------------------------------------------------------------
# Uma deque por sensor. Cada item é (horário, valor).
DADOS = {sensor["id"]: deque(maxlen=MAX_PONTOS) for sensor in SENSORES}
TRAVA = threading.Lock()

# Quanto tempo (segundos) um valor continua valendo. None = vale para sempre.
VALIDADE = {sensor["id"]: sensor.get("validade_s") for sensor in SENSORES}

# Faixa em que a leitura é possível. Fora dela, a leitura é descartada.
FAIXA_VALIDA = {sensor["id"]: sensor.get("faixa_valida") for sensor in SENSORES}

# True para o sensor cujo módulo a placa avisou que está ausente (campo opcional "max:0").
FALHA_MODULO = {}
# Última vez que a placa disse "módulo presente" (max:1). Se o sensor não manda valor mas o
# módulo está presente (ex.: MAX30102 sem dedo), o cartão fica "Sem leitura" em vez de sumir.
MODULO_VISTO = {}
MODULO_VALIDADE_S = 3

# Histórico de ocorrências (alertas e retornos ao normal). Fica só na memória.
HISTORICO = deque(maxlen=MAX_OCORRENCIAS)
TRAVA_HIST = threading.Lock()

# Situação da conexão com a Mega (a thread leitora atualiza, a página lê).
STATUS_SERIAL = {"conectado": False, "mensagem": "Procurando a placa..."}

# Conexão serial aberta (para a central ENVIAR comandos, como o do buzzer). None = sem conexão.
CONEXAO_SERIAL = None
TRAVA_ESCRITA = threading.Lock()

# Última situação do GPS (a thread leitora atualiza, a página lê).
GPS = {"estado": None, "satelites": None, "lat": None, "lon": None, "horario": None}
TRILHA = deque(maxlen=MAX_TRILHA)   # posições válidas recentes: (lat, lon)
TRAVA_GPS = threading.Lock()

# MQ-135: última leitura bruta (ADC, Rs, razão) e histórico recente da razão (para a "provável causa").
MQ135_ULTIMO = {"adc": None, "rs": None, "razao": None, "horario": None}
RAZOES_MQ135 = deque(maxlen=MAX_PONTOS)   # (horário, razão Rs/R0)


# ---------------------------------------------------------------
# GÁS (MQ-135): CÁLCULOS
# ---------------------------------------------------------------
def calcular_rs(adc):
    """ADC (0 a 1023) -> resistência do sensor em kΩ. Devolve None se o ADC for 0 ou 1023
    (sensor desligado, em curto ou saturado: não dá para calcular)."""
    if adc <= 0 or adc >= ADC_MAXIMO:
        return None
    return MQ135_RL_KOHM * (ADC_MAXIMO - adc) / adc


def obter_r0():
    """R0 em uso. No --teste vale o simulado; no modo normal vale o que você configurou (ou None)."""
    if MODO_TESTE:
        return MQ135_R0_KOHM_SIMULADO
    return MQ135_R0_KOHM


def registrar_mq135(adc, agora):
    """Recebe um valor de ADC do MQ-135 e atualiza tudo que depende dele."""
    rs = calcular_rs(adc)
    if rs is None:
        return
    r0 = obter_r0()
    razao = rs / r0 if r0 else None
    with TRAVA:
        MQ135_ULTIMO.update(adc=adc, rs=rs, razao=razao, horario=agora)
        if razao is not None:
            RAZOES_MQ135.append((agora, razao))
            DADOS["mq135"].append((agora, ppm_do_gas(MQ135_GAS_PRINCIPAL, razao)))


def avaliar_gas():
    """Avalia a última leitura do gás. Devolve (estado, info).
    estado = 'ar_normal', 'ar_atencao', 'ar_alerta' ou 'sem_dados'.
    info = dict com a razão, o ppm equivalente de cada gás e quais passaram do limite."""
    with TRAVA:
        razao, horario = MQ135_ULTIMO["razao"], MQ135_ULTIMO["horario"]
    if razao is None or horario is None:
        return "sem_dados", None
    if (datetime.now() - horario).total_seconds() > VALIDADE["mq135"]:
        return "sem_dados", None
    ppms = {g: ppm_do_gas(g, razao) for g in GASES_MQ135}
    acima = [g for g, c in GASES_MQ135.items() if ppms[g] >= c["limite_ppm"]]
    info = {"razao": razao, "ppms": ppms, "acima": acima}
    if razao <= RAZAO_ALERTA:
        return "ar_alerta", info
    if razao <= RAZAO_ATENCAO:
        return "ar_atencao", info
    return "ar_normal", info


def estado_mq135():
    if not STATUS_SERIAL["conectado"]:
        return "desconectado"
    with TRAVA:
        horario = MQ135_ULTIMO["horario"]
    if horario is None or (datetime.now() - horario).total_seconds() > VALIDADE["mq135"]:
        return "sem_dados"
    if obter_r0() is None:
        return "sem_calibracao"
    return avaliar_gas()[0]


def estimar_causa_gas():
    """PALPITE sobre o que está afetando o sensor, olhando como a razão Rs/R0 se comportou
    nos últimos segundos. Um MQ-135 sozinho NÃO identifica o gás: isto é só uma pista."""
    estado, info = avaliar_gas()
    if info is None:
        return ""
    if info["razao"] > MQ135_RAZAO_AR_LIMPO * 0.85:
        return "Provável: ar limpo, sem indício de gás"
    agora = datetime.now()
    with TRAVA:
        recentes = [(t, r) for t, r in RAZOES_MQ135 if (agora - t).total_seconds() <= 15]
    queda = 0.0
    if recentes and (agora - recentes[0][0]).total_seconds() >= 5:
        referencia = recentes[0][1]
        queda = (referencia - info["razao"]) / referencia if referencia > 0 else 0.0
    if queda >= 0.10:
        return "Provável: vapores (amônia, álcool, acetona, tolueno) • estimativa"
    return "Provável: CO₂ / ventilação ruim, ou gás já presente • estimativa"


def texto_equivalentes_gas():
    """Linha com o ppm equivalente de cada gás candidato (vai no título do gráfico grande)."""
    estado, info = avaliar_gas()
    if info is None:
        return ""
    partes = [f"{GASES_MQ135[g]['curto']} {info['ppms'][g]:.0f}" for g in GASES_MQ135]
    return "Equivalentes agora (ppm): " + " · ".join(partes)


def calibrar_r0(duracao_s=300):
    """Modo --calibrar: com o sensor em AR LIMPO e já aquecido, mede o Rs por alguns minutos
    e imprime o R0 para você copiar em MQ135_R0_KOHM."""
    threading.Thread(target=ler_serial, daemon=True).start()
    print("[calibrar] Deixe o sensor em ar limpo e ventilado (longe de álcool, perfume, fumaça).")
    print("[calibrar] O datasheet do MQ-135 recomenda aquecer por bastante tempo antes (idealmente 24 h).")
    print(f"[calibrar] Medindo por {duracao_s} s...")
    amostras, ultimo_horario = [], None
    inicio = time.time()
    while time.time() - inicio < duracao_s:
        with TRAVA:
            rs, horario = MQ135_ULTIMO["rs"], MQ135_ULTIMO["horario"]
        if rs is not None and horario != ultimo_horario:
            ultimo_horario = horario
            amostras.append(rs)
        restante = int(duracao_s - (time.time() - inicio))
        if restante % 10 == 0 and amostras:
            print(f"[calibrar] faltam {restante:3d} s | Rs atual = {amostras[-1]:.2f} kΩ | {len(amostras)} amostras")
        time.sleep(1)
    if len(amostras) < 10:
        print("[calibrar] Poucas amostras. A placa está conectada e enviando 'mq135:'? Tente de novo.")
        return
    rs_ar_limpo = statistics.median(amostras)
    r0 = rs_ar_limpo / MQ135_RAZAO_AR_LIMPO
    print()
    print(f"[calibrar] Rs em ar limpo (mediana) = {rs_ar_limpo:.2f} kΩ")
    print(f"[calibrar] R0 = Rs / {MQ135_RAZAO_AR_LIMPO} = {r0:.2f} kΩ")
    print(f"[calibrar] Copie para o topo do central_mca.py:   MQ135_R0_KOHM = {r0:.2f}")


# ---------------------------------------------------------------
# COMUNICAÇÃO SERIAL
# ---------------------------------------------------------------
def achar_porta():
    """Procura a placa pelo nome do conversor (CH340). Devolve a COM ou None."""
    for p in serial.tools.list_ports.comports():
        if "CH340" in p.description:
            return p.device
    return None


def descrever_portas():
    """Texto com as portas que o Windows está mostrando (ajuda a diagnosticar)."""
    portas = [f"{p.device} ({p.description})" for p in serial.tools.list_ports.comports()]
    return ", ".join(portas) if portas else "nenhuma porta serial encontrada"


def limpar_modulos():
    """Esquece o que a placa disse sobre os módulos (usado quando a conexão cai)."""
    FALHA_MODULO.clear()
    MODULO_VISTO.clear()


def definir_status(conectado, mensagem, detalhe=None):
    """Atualiza o status da conexão e escreve no terminal quando a mensagem MUDA."""
    mudou = STATUS_SERIAL["mensagem"] != mensagem
    STATUS_SERIAL.update(conectado=conectado, mensagem=mensagem)
    if mudou:
        print("[serial]", mensagem)
        if detalhe:
            print("[serial]  ->", detalhe)


def enviar_comando(texto):
    """Manda uma linha de texto para a Mega (ex.: 'buz:1'). Devolve True se conseguiu enviar."""
    conexao = CONEXAO_SERIAL
    if conexao is None:
        return False
    try:
        with TRAVA_ESCRITA:
            conexao.write((texto + "\n").encode("ascii"))
        return True
    except (serial.SerialException, OSError):
        return False


def interpretar_linha(linha):
    """Converte 'mq135:512,lux:180.0,bpm:72,spo2:97' em {'mq135': 512.0, ...}.
    Ignora o que não entender, valores que não são número e leituras fora da faixa
    em que o sensor consegue medir (ex.: DS18B20 desconectado devolve -127)."""
    resultado = {}
    for parte in linha.split(","):
        if ":" not in parte:
            continue
        nome, texto = parte.split(":", 1)
        sensor_id = MAPA_SERIAL.get(nome.strip())
        if sensor_id is None:
            continue  # nome que não está no mapa
        try:
            valor = float(texto)
        except ValueError:
            continue  # valor que não é número
        if not math.isfinite(valor):
            continue  # "nan" ou "inf"
        faixa = FAIXA_VALIDA.get(sensor_id)
        if faixa and not (faixa[0] <= valor <= faixa[1]):
            continue  # impossível para esse sensor: descarta
        resultado[sensor_id] = valor
    return resultado


def interpretar_modulos(linha):
    """Lê campos opcionais como 'max:1' / 'max:0' (módulo presente / ausente)."""
    resultado = {}
    for parte in linha.split(","):
        if ":" not in parte:
            continue
        nome, texto = parte.split(":", 1)
        nome = nome.strip()
        if nome in MODULOS_SERIAL:
            try:
                resultado[nome] = float(texto) != 0
            except ValueError:
                continue
    return resultado


def interpretar_gps(linha):
    """Extrai só os campos do GPS de uma linha como
    'mq135:38,gps:2,sat:7,lat:-22.1,lon:-43.1'. Ignora o que não entender."""
    campos = {}
    for parte in linha.split(","):
        if ":" not in parte:
            continue
        nome, texto = parte.split(":", 1)
        nome = nome.strip()
        if nome in ("gps", "sat", "lat", "lon"):
            try:
                campos[nome] = float(texto)
            except ValueError:
                continue
    return campos


def atualizar_gps(linha):
    """Atualiza a situação do GPS. Só age em linhas que trazem o campo 'gps'."""
    campos = interpretar_gps(linha)
    if "gps" not in campos:
        return  # linha de erro, linha cortada ou placa sem GPS
    lat, lon = campos.get("lat"), campos.get("lon")
    # Posição só vale se o módulo disse 2 E as duas coordenadas existem e fazem sentido
    posicao_valida = (
        campos["gps"] == 2
        and lat is not None and lon is not None
        and -90 <= lat <= 90 and -180 <= lon <= 180
    )
    with TRAVA_GPS:
        GPS["estado"] = int(campos["gps"])
        GPS["satelites"] = int(campos["sat"]) if "sat" in campos else None
        GPS["lat"] = lat if posicao_valida else None   # sem posição: apaga a antiga
        GPS["lon"] = lon if posicao_valida else None
        GPS["horario"] = datetime.now()
        if posicao_valida:
            # Guarda apenas posições novas, evitando repetir o mesmo ponto no trajeto.
            nova_posicao = (lat, lon)
            ultima_posicao = TRILHA[-1] if TRILHA else None
            if ultima_posicao is None or abs(ultima_posicao[0] - lat) > 0.00001 or abs(ultima_posicao[1] - lon) > 0.00001:
                TRILHA.append(nova_posicao)


def processar_linha(linha):
    """Recebe UMA linha da placa (real ou simulada) e distribui para GPS, módulos e sensores.
    É o único caminho de entrada de dados: a placa real e o modo teste passam por aqui."""
    atualizar_gps(linha)
    for modulo, presente in interpretar_modulos(linha).items():
        for sensor_id in MODULOS_SERIAL[modulo]:
            FALHA_MODULO[sensor_id] = not presente
            if presente:
                MODULO_VISTO[sensor_id] = datetime.now()
    leituras = interpretar_linha(linha)
    if not leituras:
        return  # linha vazia ou inválida
    agora = datetime.now()
    # O MQ-135 chega como ADC bruto: tem um caminho próprio (ADC -> Rs -> razão -> ppm estimado).
    if "mq135" in leituras:
        registrar_mq135(leituras.pop("mq135"), agora)
    with TRAVA:
        for sensor_id, valor in leituras.items():
            DADOS[sensor_id].append((agora, valor))


def ler_serial():
    """Lê a Mega pela serial. Se a placa sumir, avisa e tenta de novo."""
    global CONEXAO_SERIAL
    while True:
        porta = PORTA_MANUAL or achar_porta()
        if porta is None:
            definir_status(False, "Placa (CH340) não encontrada. Confira o cabo USB.",
                           detalhe="Portas vistas pelo Windows: " + descrever_portas())
            limpar_modulos()
            time.sleep(2)
            continue
        try:
            with serial.Serial(porta, BAUD, timeout=2) as conexao:
                CONEXAO_SERIAL = conexao
                definir_status(True, f"Conectado em {porta} a {BAUD} baud")
                try:
                    while True:
                        linha = conexao.readline().decode("utf-8", errors="ignore").strip()
                        processar_linha(linha)
                        if DIAGNOSTICO_SERIAL and linha:
                            print("RECEBIDO:", repr(linha))
                            print("   GPS lido:", interpretar_gps(linha), "| estado do GPS:", estado_do_gps())
                finally:
                    CONEXAO_SERIAL = None
        except (serial.SerialException, OSError) as erro:
            # Também cai aqui se outro programa (ex.: Monitor Serial do Arduino IDE) estiver usando a porta.
            definir_status(False, f"Falha na porta: {erro}")
            limpar_modulos()
            time.sleep(2)


# ---------------------------------------------------------------
# MODO TESTE: PLACA SIMULADA (só com --teste)
# ---------------------------------------------------------------
def montar_linha_simulada(t):
    """Monta a linha que a placa enviaria no segundo t, seguindo um roteiro de 120 s
    que força cada situação (dedo fora, sensor sumindo, GPS sem fix, alerta de gás...).
    DADOS FALSOS, só para testar a interface."""
    c = t % 120
    campos = []

    # Com R0 simulado = 20 kΩ e RL = 10 kΩ: ADC ~125 = ar limpo (razão ~3,6, poucos ppm);
    # ADC ~190 = ATENÇÃO; ADC ~255 = ALERTA (razão ~1,5). Só o roteiro mexe nesses números.
    if 100 <= c < 115:
        adc = 255 + random.gauss(0, 3)
    elif 95 <= c < 100:
        adc = 190 + random.gauss(0, 3)
    else:
        adc = 125 + 6 * math.sin(t / 15) + random.gauss(0, 2)
    campos.append(f"mq135:{adc:.0f}")

    if not (90 <= c < 105):                       # 90-105 s: luminosidade para de enviar
        campos.append(f"lux:{300 + 30 * math.sin(t / 10):.1f}")

    if 80 <= c < 95:                              # 80-95 s: DS18B20 "desconectado" (-127)
        campos.append("temp:-127.00")
    else:
        campos.append(f"temp:{27 + 0.5 * math.sin(t / 20):.1f}")

    if c < 30 or c >= 65:                         # dedo no MAX30102
        campos += [f"bpm:{72 + 3 * math.sin(t / 5):.0f}", f"spo2:{97 + math.sin(t / 7):.0f}", "max:1"]
    elif c < 45:                                  # 30-45 s: sem dedo (como o firmware real: omite os campos)
        campos += ["max:1"]
    elif c < 55:                                  # 45-55 s: placa para de enviar BPM, SpO2 e max
        pass
    else:                                         # 55-65 s: MAX30102 desconectado (max:0)
        campos += ["max:0"]

    if c < 20:                                    # 0-20 s: GPS procurando satélites
        campos += ["gps:1", "sat:3"]
    elif c < 100:                                 # 20-100 s: GPS com posição
        passo = 0.00003 * (c - 20)
        campos += ["gps:2", "sat:7", f"lat:{-22.9068 + passo:.6f}", f"lon:{-43.1729 + passo:.6f}"]
    # 100-120 s: GPS para de enviar
    return ",".join(campos)


def simular_placa():
    """Thread do modo teste: alimenta a central com linhas falsas, uma por segundo."""
    definir_status(True, "MODO TESTE: placa simulada (dados falsos)")
    inicio = time.time()
    while True:
        processar_linha(montar_linha_simulada(time.time() - inicio))
        time.sleep(1)


def simular_mq135():
    """DEMONSTRAÇÃO: gera valores de ADC falsos de ar limpo para o MQ-135.
    Só roda se algum sensor do cadastro tiver origem 'simulado'."""
    inicio = time.time()
    while True:
        t = time.time() - inicio
        registrar_mq135(125 + 6 * math.sin(t / 15) + random.gauss(0, 2), datetime.now())
        time.sleep(INTERVALO_S)


# ---------------------------------------------------------------
# LEITURA DE VALORES E ESTADOS
# ---------------------------------------------------------------
def ultimo_valor(sensor_id):
    """Devolve o valor mais recente do sensor, ou None se não há dados
    (ou se o último valor é velho demais, conforme 'validade_s')."""
    with TRAVA:
        dados = DADOS[sensor_id]
        if not dados:
            return None
        horario, valor = dados[-1]
    validade = VALIDADE[sensor_id]
    if validade is not None and (datetime.now() - horario).total_seconds() > validade:
        return None
    return valor


def copiar_dados(sensor_id):
    """Devolve duas listas (horários, valores) copiadas com segurança."""
    with TRAVA:
        pontos = list(DADOS[sensor_id])
    return [p[0] for p in pontos], [p[1] for p in pontos]


def avaliar_limite(sensor, valor):
    """Compara o valor com os limites do sensor.
    Devolve 'baixo', 'alto', 'ok' ou None (se o sensor não tem limites)."""
    limites = sensor.get("limites")
    if not limites or valor is None:
        return None
    minimo, maximo = limites.get("min"), limites.get("max")
    if minimo is not None and valor < minimo:
        return "baixo"
    if maximo is not None and valor > maximo:
        return "alto"
    return "ok"


def estado_do_gps():
    """Decide em qual das situações o GPS está."""
    if not STATUS_SERIAL["conectado"]:
        return "desconectado"
    with TRAVA_GPS:
        estado, horario, lat = GPS["estado"], GPS["horario"], GPS["lat"]
    if horario is None or (datetime.now() - horario).total_seconds() > VALIDADE_GPS_S:
        return "sem_dados"
    if estado == 0:
        return "gps_sem_modulo"
    if estado == 2 and lat is not None:
        return "gps_ok"
    return "gps_sem_sinal"


def modulo_presente(sensor_id):
    """True se a placa avisou 'módulo presente' (max:1) nos últimos MODULO_VALIDADE_S segundos."""
    visto = MODULO_VISTO.get(sensor_id)
    return (visto is not None and not FALHA_MODULO.get(sensor_id)
            and (datetime.now() - visto).total_seconds() <= MODULO_VALIDADE_S)


def estado_do_sensor(sensor):
    """Decide o estado do cartão com base na origem, nos dados e nos limites."""
    if sensor["origem"] == "gps":
        return estado_do_gps()
    if sensor["id"] == "mq135":
        return estado_mq135()
    if sensor["origem"] == "serial" and not STATUS_SERIAL["conectado"]:
        return "desconectado"
    if FALHA_MODULO.get(sensor["id"]):
        return "sem_dados"            # a placa avisou que o módulo (ex.: MAX30102) não responde
    valor = ultimo_valor(sensor["id"])
    ocioso = sensor.get("ocioso")
    if valor is None:
        # Sem valor: se a placa disse há pouco que o módulo está presente, é "sem dedo / medindo".
        if ocioso and modulo_presente(sensor["id"]):
            return "ocioso"
        return "sem_dados"            # parou de chegar: sai da Conexão (mesmo BPM e SpO2)
    if ocioso and valor <= ocioso["valor"]:
        return "ocioso"               # a placa está falando, mas sem leitura (ex.: sem dedo)
    if sensor["origem"] == "simulado":
        return "simulado"
    resultado = avaliar_limite(sensor, valor)
    if resultado == "baixo":
        return "alerta_baixo"
    if resultado == "alto":
        return "alerta_alto"
    if resultado == "ok":
        return "normal"
    return "real"  # sensor sem limites cadastrados


def sensor_ativo(sensor):
    """True se o sensor está enviando dados agora."""
    return estado_do_sensor(sensor) not in ESTADOS_INATIVOS


def texto_estado(sensor, estado):
    """Texto pequeno do cartão."""
    return ESTADOS[estado]["texto"]


def texto_gps():
    """Texto do cartão do GPS (com quebras de linha)."""
    estado = estado_do_gps()
    with TRAVA_GPS:
        lat, lon, sat = GPS["lat"], GPS["lon"], GPS["satelites"]
    if estado == "gps_ok" and lat is not None:
        return f"Lat {lat:.5f}\nLon {lon:.5f}\n{sat or 0} satélites"
    if estado == "gps_sem_sinal":
        return f"--\n{sat or 0} satélites"
    return "--"


def texto_valor(sensor):
    """Texto grande do cartão, como '352 ' ou '-- °C'."""
    if sensor["origem"] == "gps":
        return texto_gps()
    if estado_do_sensor(sensor) in ("desconectado", "sem_dados"):
        return f"-- {sensor['unidade']}"
    valor = ultimo_valor(sensor["id"])
    if valor is None:
        return f"-- {sensor['unidade']}"
    return f"{valor:.{sensor['casas']}f} {sensor['unidade']}"


def frase_alerta_gas(estado):
    """Frase da faixa de alertas para o gás (ou '' se não houver como montar)."""
    valor = ultimo_valor("mq135")
    _, info = avaliar_gas()
    if valor is None or info is None:
        return ""
    if estado == "ar_alerta":
        nomes = ", ".join(GASES_MQ135[g]["curto"] for g in info["acima"]) or "algum gás"
        return (f"Gás: {valor:.1f} ppm eq. NH₃, ACIMA do limite de pior caso (NR-15). "
                f"Equivalente acima do limite para: {nomes}. O MQ-135 não identifica o gás: trate como suspeita.")
    return (f"Gás: {valor:.1f} ppm eq. NH₃, acima de {MQ135_FRACAO_ATENCAO:.0%} do limite de pior caso (atenção). "
            "Estimativa.")


def lista_de_alertas():
    """Devolve as frases de todos os sensores que estão fora do limite agora."""
    frases = []
    for sensor in SENSORES:
        estado = estado_do_sensor(sensor)
        if estado not in ESTADOS_DE_ALERTA:
            continue
        if sensor["id"] == "mq135":
            frase = frase_alerta_gas(estado)
            if frase:
                frases.append(frase)
            continue
        limites = sensor["limites"]
        if estado == "alerta_baixo":
            detalhe = f"abaixo do mínimo configurado ({limites['min']} {sensor['unidade']})"
        else:
            detalhe = f"acima do máximo configurado ({limites['max']} {sensor['unidade']})"
        frases.append(f"{sensor['nome']}: {texto_valor(sensor).strip()}, {detalhe}")
    return frases


# ---------------------------------------------------------------
# HISTÓRICO DE OCORRÊNCIAS
# ---------------------------------------------------------------
def registrar_ocorrencia(sensor, estado):
    """Guarda uma linha no histórico. 'estado' é um alerta, 'normal', 'ativo' ou um estado inativo."""
    if estado == "normal":
        texto = "Voltou ao normal"
    elif estado == "ativo":
        texto = "Sensor ativo"
    else:
        texto = ESTADOS[estado]["texto"]
    ocorrencia = {
        "horario": datetime.now(),
        "sensor": sensor["nome"],
        "texto": texto,
        "valor": f"{texto_valor(sensor).strip()}",
        "alerta": estado in ESTADOS_DE_ALERTA,
    }
    with TRAVA_HIST:
        HISTORICO.append(ocorrencia)


def verificar_mudancas(anterior):
    """Uma passada: compara o estado de cada sensor com o anterior e registra só quando MUDA.
    'anterior' é um dicionário {id: estado} que a função atualiza."""
    for sensor in SENSORES:
        estado = estado_do_sensor(sensor)
        antes = anterior.get(sensor["id"])
        if estado in ESTADOS_DE_ALERTA and estado != antes:
            registrar_ocorrencia(sensor, estado)
        elif estado in ESTADOS_NORMAIS and antes in ESTADOS_DE_ALERTA:
            registrar_ocorrencia(sensor, "normal")
        elif estado in ESTADOS_INATIVOS and antes is not None and antes not in ESTADOS_INATIVOS:
            registrar_ocorrencia(sensor, estado)      # o sensor parou de enviar dados
        elif estado not in ESTADOS_INATIVOS and antes in ESTADOS_INATIVOS:
            registrar_ocorrencia(sensor, "ativo")     # o sensor voltou a enviar dados
        anterior[sensor["id"]] = estado


def vigiar_alertas():
    """Thread de vigilância: registra mudanças de estado e, uma vez por segundo, avisa a Mega
    se o gás está em risco ('buz:1') ou normal ('buz:0'). A Mega desliga o alarme sozinha se
    ficar alguns segundos sem receber o comando (por exemplo, se a central fechar)."""
    anterior = {}
    ultimo_comando = 0.0
    while True:
        verificar_mudancas(anterior)
        if time.time() - ultimo_comando >= 1.0:
            ultimo_comando = time.time()
            enviar_comando("buz:1" if estado_mq135() == "ar_alerta" else "buz:0")
        time.sleep(0.5)


def historico_para_csv(ocorrencias):
    """Texto CSV do histórico, pronto para o Excel em português (separador ';' e acentos corretos)."""
    saida = io.StringIO()
    escritor = csv.writer(saida, delimiter=";", lineterminator="\r\n")
    escritor.writerow(["Data", "Hora", "Sensor", "Ocorrência", "Valor", "Severidade"])
    for o in ocorrencias:
        escritor.writerow([
            o["horario"].strftime("%d/%m/%Y"), o["horario"].strftime("%H:%M:%S"),
            o["sensor"], o["texto"], o["valor"], "Alerta" if o["alerta"] else "Informativo",
        ])
    return "\ufeff" + saida.getvalue()   # \ufeff = marca que faz o Excel reconhecer os acentos


def historico_para_xlsx(ocorrencias):
    """Gera um arquivo Excel (.xlsx) em bytes com duas planilhas:
    'Ocorrências' (tabela completa, com filtro) e 'Resumo' (totais + gráfico por sensor)."""
    wb = Workbook()

    # ---------- Planilha 1: Ocorrências ----------
    ws = wb.active
    ws.title = "Ocorrências"
    ws.append(["Data/Hora", "Sensor", "Ocorrência", "Valor", "Severidade"])

    fonte_cab = Font(bold=True, color="FFFFFF")
    fundo_cab = PatternFill("solid", fgColor="262D38")
    for celula in ws[1]:
        celula.font = fonte_cab
        celula.fill = fundo_cab
        celula.alignment = Alignment(horizontal="center")

    fundo_alerta = PatternFill("solid", fgColor="FADBD8")   # vermelho claro
    for o in ocorrencias:
        ws.append([o["horario"], o["sensor"], o["texto"], o["valor"],
                   "Alerta" if o["alerta"] else "Informativo"])
        linha = ws.max_row
        ws.cell(linha, 1).number_format = "dd/mm/yyyy hh:mm:ss"   # data de verdade, não texto
        if o["alerta"]:
            for coluna in range(1, 6):
                ws.cell(linha, coluna).fill = fundo_alerta

    for coluna, largura in enumerate((20, 24, 38, 18, 14), start=1):
        ws.column_dimensions[get_column_letter(coluna)].width = largura
    ws.freeze_panes = "A2"                       # cabeçalho fixo ao rolar
    ws.auto_filter.ref = ws.dimensions           # setinhas de filtro no cabeçalho

    # ---------- Planilha 2: Resumo ----------
    resumo = wb.create_sheet("Resumo")
    total = len(ocorrencias)
    alertas = sum(1 for o in ocorrencias if o["alerta"])
    resumo.append(["Resumo do histórico"])
    resumo["A1"].font = Font(bold=True, size=14)
    resumo.append(["Gerado em", datetime.now()])
    resumo["B2"].number_format = "dd/mm/yyyy hh:mm:ss"
    resumo.append(["Total de ocorrências", total])
    resumo.append(["Alertas", alertas])
    resumo.append(["Informativos", total - alertas])
    if ocorrencias:
        resumo.append(["Primeira ocorrência", ocorrencias[0]["horario"]])
        resumo.append(["Última ocorrência", ocorrencias[-1]["horario"]])
        resumo["B6"].number_format = resumo["B7"].number_format = "dd/mm/yyyy hh:mm:ss"

    # Tabela por sensor
    contagem = {}
    for o in ocorrencias:
        item = contagem.setdefault(o["sensor"], [0, 0])
        item[0 if o["alerta"] else 1] += 1
    inicio = resumo.max_row + 2
    resumo.cell(inicio, 1, "Sensor")
    resumo.cell(inicio, 2, "Alertas")
    resumo.cell(inicio, 3, "Informativos")
    for celula in resumo[inicio]:
        celula.font = fonte_cab
        celula.fill = fundo_cab
    for sensor_nome, (n_alertas, n_info) in sorted(contagem.items()):
        resumo.append([sensor_nome, n_alertas, n_info])

    if contagem:
        grafico = BarChart()
        grafico.type = "col"
        grafico.grouping = "stacked"
        grafico.overlap = 100
        grafico.title = "Ocorrências por sensor"
        dados = Reference(resumo, min_col=2, max_col=3, min_row=inicio, max_row=resumo.max_row)
        nomes = Reference(resumo, min_col=1, min_row=inicio + 1, max_row=resumo.max_row)
        grafico.add_data(dados, titles_from_data=True)
        grafico.set_categories(nomes)
        grafico.height, grafico.width = 8, 16
        resumo.add_chart(grafico, "E2")

    resumo.column_dimensions["A"].width = 24
    resumo.column_dimensions["B"].width = 22
    resumo.column_dimensions["C"].width = 14

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def iniciar_threads():
    """Liga as threads de fundo. Chamada só quando o arquivo é executado diretamente."""
    if MODO_TESTE:
        threading.Thread(target=simular_placa, daemon=True).start()
    elif any(s["origem"] in ("serial", "gps") for s in SENSORES):
        threading.Thread(target=ler_serial, daemon=True).start()
    if any(s["origem"] == "simulado" for s in SENSORES):
        threading.Thread(target=simular_mq135, daemon=True).start()
    threading.Thread(target=vigiar_alertas, daemon=True).start()


# ---------------------------------------------------------------
# COMPONENTES VISUAIS
# ---------------------------------------------------------------
COR_FUNDO = "#0E1116"
COR_SUPERFICIE = "#151A21"
COR_SUPERFICIE_2 = "#1C222B"
COR_BORDA = "#262D38"
COR_TEXTO = "#E6E8EB"
COR_SECUNDARIA = "#8B95A3"
COR_DESTAQUE = "#D98B55"   # cobre


def copiar_historico():
    with TRAVA_HIST:
        return list(HISTORICO)


def estilo_grafico_historico(fig, titulo, uirevision):
    """Aplica o visual do painel (fundo transparente, grade discreta) a um gráfico do histórico."""
    fig.update_layout(
        title=dict(text=titulo, x=0, xanchor="left", font=dict(size=14, color=COR_SECUNDARIA)),
        height=250, margin=dict(l=40, r=16, t=40, b=50),
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
        font=dict(color=COR_SECUNDARIA, family="Segoe UI"),
        legend=dict(orientation="h", y=-0.3), hovermode="x unified", uirevision=uirevision,
    )
    fig.update_xaxes(gridcolor=COR_BORDA, zeroline=False)
    fig.update_yaxes(gridcolor=COR_BORDA, zeroline=False, rangemode="tozero")
    return fig


def criar_graficos_historico():
    """Dois gráficos feitos só com as ocorrências registradas nesta sessão:
    barras por hora (hoje) e linhas dos últimos 7 dias. Duas categorias: alerta e informativo."""
    ocorrencias = copiar_historico()
    hoje = datetime.now().date()
    dias = [datetime.fromordinal(hoje.toordinal() - i).date() for i in range(6, -1, -1)]

    alertas_hora, info_hora = [0] * 24, [0] * 24
    alertas_dia, info_dia = [0] * 7, [0] * 7
    for o in ocorrencias:
        quando = o["horario"]
        if quando.date() == hoje:
            (alertas_hora if o["alerta"] else info_hora)[quando.hour] += 1
        if quando.date() in dias:
            (alertas_dia if o["alerta"] else info_dia)[dias.index(quando.date())] += 1

    horas = list(range(24))
    fig_hora = go.Figure()
    fig_hora.add_bar(x=horas, y=alertas_hora, name="Alertas", marker_color="#F0616D")
    fig_hora.add_bar(x=horas, y=info_hora, name="Informativos", marker_color=COR_DESTAQUE)
    fig_hora.update_layout(barmode="stack")
    estilo_grafico_historico(fig_hora, "Ocorrências por hora • hoje", "hist-horas")
    fig_hora.update_xaxes(tickmode="linear", dtick=2)

    rotulos = [d.strftime("%d/%m") for d in dias]
    fig_dias = go.Figure()
    fig_dias.add_scatter(x=rotulos, y=alertas_dia, name="Alertas", mode="lines+markers",
                         line=dict(color="#F0616D", width=2), marker=dict(size=6))
    fig_dias.add_scatter(x=rotulos, y=info_dia, name="Informativos", mode="lines+markers",
                         line=dict(color=COR_DESTAQUE, width=2), marker=dict(size=6))
    estilo_grafico_historico(fig_dias, "Ocorrências • últimos 7 dias", "hist-dias")

    return fig_hora, fig_dias


def criar_tabela_historico():
    """Tabela de eventos recentes, do mais novo ao mais antigo. O estilo vem do CSS (.tabela, .sev)."""
    ocorrencias = copiar_historico()
    if not ocorrencias:
        return html.Div("Nenhuma ocorrência registrada ainda. Os eventos aparecem aqui quando um sensor muda de estado.",
                        className="vazio")
    cabecalho = html.Tr([html.Th(t) for t in ("Horário", "Sensor", "Ocorrência", "Valor", "Severidade")])
    linhas = []
    for o in reversed(ocorrencias):
        classe_sev = "sev sev-alerta" if o["alerta"] else "sev sev-info"
        linhas.append(html.Tr([
            html.Td(o["horario"].strftime("%d/%m/%Y %H:%M:%S"), className="num"),
            html.Td(o["sensor"]),
            html.Td(o["texto"]),
            html.Td(o["valor"], className="num"),
            html.Td(html.Span("Alerta" if o["alerta"] else "Informativo", className=classe_sev)),
        ]))
    return html.Table([html.Thead(cabecalho), html.Tbody(linhas)], className="tabela")


def criar_indicador(titulo, valor, id_valor, classe=""):
    """Cartão de número do histórico. O id é o mesmo que o callback atualizar_tela_historico usa."""
    return html.Div([html.Div(titulo, className="ind-titulo"),
                     html.Div(str(valor), id=id_valor, className=f"ind-valor {classe}".strip())],
                    className="indicador")


def criar_tela_historico():
    """Tela completa da aba Histórico de eventos:
    gráficos empilhados na coluna da esquerda, registro de eventos na direita."""
    ocorrencias = copiar_historico()
    total = len(ocorrencias)
    alertas = sum(1 for o in ocorrencias if o["alerta"])
    fig_hora, fig_dias = criar_graficos_historico()
    return html.Div([
        html.Div([html.H1("Histórico de eventos"),
                  html.Div("Alertas e mudanças de estado desta sessão", className="hora")],
                 className="titulo-linha"),
        html.Div([
            criar_indicador("Ocorrências", total, "historico-total"),
            criar_indicador("Alertas", alertas, "historico-alertas-total", "alerta"),
            criar_indicador("Informativos", total - alertas, "historico-info-total", "cobre"),
        ], className="grade-indicadores"),
        html.Div([
            html.Div([
                html.Section(dcc.Graph(id="historico-grafico-horas", figure=fig_hora, config={"displayModeBar": False}), className="painel"),
                html.Section(dcc.Graph(id="historico-grafico-dias", figure=fig_dias, config={"displayModeBar": False}), className="painel"),
            ], className="h-graficos"),
            html.Section([
                html.Div([
                    html.Div("Registro detalhado", className="rotulo"),
                    html.Button("Salvar como Excel (.xlsx)", id="botao-salvar-excel", n_clicks=0, className="botao-acao"),
                ], className="h-registro-topo"),
                html.Div(criar_tabela_historico(), id="tabela-historico-aba", className="tabela-wrap h-tabela-rolavel"),
                html.Div("O histórico fica só na memória: ao fechar a central ele é perdido. Salve antes de encerrar.",
                         id="aviso-salvar", className="nota"),
                dcc.Download(id="download-historico"),
            ], className="painel h-registro"),
        ], className="h-grade"),
    ])


def texto_resumo_gps():
    """Frase da faixa no topo da aba GPS, de acordo com a situação real do módulo."""
    estado = estado_do_gps()
    with TRAVA_GPS:
        lat, lon, sat = GPS["lat"], GPS["lon"], GPS["satelites"]
    if estado == "gps_ok" and lat is not None:
        return f"Latitude {lat:.6f} • Longitude {lon:.6f} • {sat or 0} satélites"
    if estado == "gps_sem_sinal":
        return f"Sem posição: o módulo vê {sat or 0} satélites. Precisa de céu aberto."
    if estado == "gps_sem_modulo":
        return "A placa informa que o módulo GPS não está respondendo."
    if estado == "desconectado":
        return "Placa desconectada."
    return f"Nenhum dado de GPS chegou nos últimos {VALIDADE_GPS_S} segundos."


def criar_tela_gps():
    """Tela completa da aba GPS / Localização."""
    return html.Div([
        html.Div([html.H1("GPS / Localização"),
                  html.Div("Posição atual e trajeto recente", className="hora")],
                 className="titulo-linha"),
        html.Div(texto_resumo_gps(), id="gps-resumo", className="faixa-alerta ok"),
        html.Section(dcc.Graph(id="mapa-gps-separado", figure=criar_mapa(), className="mapa",
                               config={"displayModeBar": True, "scrollZoom": True, "displaylogo": False}),
                     className="painel"),
        html.P("O mapa usa imagens da internet (OpenStreetMap): sem internet, a posição é válida mas o fundo do mapa não aparece.",
               className="nota"),
        html.P("Para vários coletes simultâneos, cada dispositivo precisará enviar um identificador único junto com a latitude e a longitude.",
               className="nota"),
    ])


def classe_cartao(estado, aberto=False):
    """Monta o class do cartão. Sensor sem dados recebe 'oculto' e some do painel."""
    classes = ["cartao", "est-" + CATEGORIA.get(estado, "neutro")]
    if aberto:
        classes.append("aberto")
    if estado in ESTADOS_INATIVOS:
        classes.append("oculto")
    return " ".join(classes)


def criar_spark(sensor):
    """Minigráfico dos últimos pontos, sem eixos, para dentro do cartão."""
    x, y = copiar_dados(sensor["id"])
    fig = go.Figure(go.Scatter(x=x[-60:], y=y[-60:], mode="lines",
                               line=dict(color=COR_DESTAQUE, width=1.5)))
    fig.update_layout(height=48, margin=dict(l=0, r=0, t=0, b=0), showlegend=False,
                      paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
                      xaxis=dict(visible=False), yaxis=dict(visible=False))
    return fig


# ---- Informações extras do cartão: mín/méd/máx, limites e (no gás) a provável causa ----
def estatisticas(sensor_id):
    """(mínimo, média, máximo) dos pontos guardados do sensor, ou None se não há pontos."""
    _, y = copiar_dados(sensor_id)
    if not y:
        return None
    return min(y), sum(y) / len(y), max(y)


def texto_estatisticas(sensor):
    est = estatisticas(sensor["id"])
    if est is None:
        return "mín -- · méd -- · máx --"
    c = sensor["casas"]
    return f"mín {est[0]:.{c}f} · méd {est[1]:.{c}f} · máx {est[2]:.{c}f}"


def texto_limites(sensor):
    """Frase com os limites configurados. Vazio se o sensor não tem limites."""
    if sensor["id"] == "mq135":
        return (f"Atenção ≥ {ATENCAO_GAS_PPM:g} · Alerta ≥ {LIMITE_GAS_PPM:g} ppm eq. NH₃ (NR-15, pior caso)")
    limites = sensor.get("limites") or {}
    minimo, maximo = limites.get("min"), limites.get("max")
    un = sensor["unidade"]
    if minimo is not None and maximo is not None:
        return f"Limite: {minimo:g} – {maximo:g} {un}"
    if minimo is not None:
        return f"Limite: ≥ {minimo:g} {un}"
    if maximo is not None:
        return f"Limite: ≤ {maximo:g} {un}"
    return ""


def criar_extra(sensor):
    """Linhas pequenas abaixo do estado do cartão. O estilo vem do CSS (.extra-linha...)."""
    linhas = [html.Div(texto_estatisticas(sensor), className="extra-linha")]
    limite = texto_limites(sensor)
    if limite:
        linhas.append(html.Div(limite, className="extra-linha extra-limite"))
    if sensor["id"] == "mq135":
        causa = estimar_causa_gas()
        if causa:
            linhas.append(html.Div(causa, className="extra-linha extra-causa"))
    return linhas


def criar_cartao(sensor, aberto=False):
    """Cartão de sensor (clicável): título, número grande, unidade, estado, extras e minigráfico."""
    estado = estado_do_sensor(sensor)
    numero, _, unidade = texto_valor(sensor).rpartition(" ")
    selo = html.Span("EST.", className="selo", title="Valor estimado, não calibrado") if sensor["id"] in ESTIMADOS else None
    return html.Div(
        [
            html.Div([html.Span(sensor["nome"]), selo, html.Span(className="ponto")], className="cartao-topo"),
            html.Div([html.Span(numero), html.Span(unidade, className="unidade")],
                     id={"type": "valor", "id": sensor["id"]}, className="cartao-valor"),
            html.Div(texto_estado(sensor, estado), id={"type": "estado", "id": sensor["id"]}, className="cartao-estado"),
            html.Div(criar_extra(sensor), id={"type": "extra", "id": sensor["id"]}, className="cartao-extra"),
            dcc.Graph(id={"type": "spark", "id": sensor["id"]}, figure=criar_spark(sensor),
                      config={"staticPlot": True}, className="spark", style={"height": "48px"}),
        ],
        id={"type": "cartao", "id": sensor["id"]},
        n_clicks=0,
        className=classe_cartao(estado, aberto),
        title=sensor["descricao"],   # a descrição longa aparece ao passar o mouse
    )


def criar_mapa():
    """Mapa do GPS: posição atual (ponto vermelho) e trajeto recente (linha âmbar).
    Sem posição válida, mostra só um aviso (nunca uma posição antiga)."""
    with TRAVA_GPS:
        lat, lon = GPS["lat"], GPS["lon"]
        trilha = list(TRILHA)
    posicao_ok = estado_do_gps() == "gps_ok" and lat is not None
    fig = go.Figure()

    if posicao_ok:
        fig.add_trace(go.Scattermap(
            lat=[p[0] for p in trilha], lon=[p[1] for p in trilha],
            mode="lines", line=dict(width=3, color="#D4AF37"), name="Trajeto",
        ))
        fig.add_trace(go.Scattermap(
            lat=[lat], lon=[lon], mode="markers",
            marker=dict(size=14, color="#fb7185"), name="Posição atual",
        ))
        fig.update_layout(
            title=f"GPS: Lat {lat:.6f}, Lon {lon:.6f}",
            map=dict(style="open-street-map", center=dict(lat=lat, lon=lon), zoom=16),
            margin=dict(l=0, r=0, t=40, b=0),
        )
    else:
        fig.update_layout(title="GPS", template="plotly_dark")
        fig.update_xaxes(visible=False)
        fig.update_yaxes(visible=False)
        fig.add_annotation(
            text="Sem posição válida", x=0.5, y=0.5, xref="paper", yref="paper",
            showarrow=False, font=dict(size=20, color=COR_SECUNDARIA),
        )

    # O uirevision muda entre "mapa" e "sem posição": ao ganhar o primeiro fix, o mapa
    # se centraliza; depois disso, zoom e arraste do usuário são preservados.
    fig.update_layout(height=650, autosize=True,
                      uirevision="gps-mapa" if posicao_ok else "gps-sem-posicao")
    return fig


def faixa_segura(sensor, y):
    """(inferior, superior) da zona segura para sombrear no gráfico, ou None se não há limites."""
    limites = sensor.get("limites") or {}
    minimo, maximo = limites.get("min"), limites.get("max")
    if minimo is None and maximo is None:
        return None
    base = sensor.get("faixa")
    inferior = minimo if minimo is not None else (base[0] if base else 0)
    superior = maximo if maximo is not None else (base[1] if base else max(y) * 1.2)
    return inferior, superior


def criar_figura(sensor):
    """Gráfico grande do sensor: zona segura sombreada, linhas de limite, pontos fora do limite
    em vermelho e o valor atual destacado."""
    if sensor["id"] == "gps":
        return criar_mapa()
    x, y = copiar_dados(sensor["id"])
    tem_dados = len(x) > 0

    titulo = sensor["nome"]
    if tem_dados and sensor["origem"] == "simulado":
        titulo += " (DEMONSTRAÇÃO)"
    if sensor["id"] == "mq135":
        equivalentes = texto_equivalentes_gas()
        if equivalentes:
            titulo += f"<br><sup>{equivalentes}</sup>"

    fig = go.Figure()
    if tem_dados:
        limites = sensor.get("limites") or {}
        minimo, maximo = limites.get("min"), limites.get("max")

        zona = faixa_segura(sensor, y)
        if zona:
            fig.add_hrect(y0=zona[0], y1=zona[1], fillcolor="rgba(52,211,153,0.07)",
                          line_width=0, layer="below")

        fig.add_trace(go.Scatter(x=x, y=y, mode="lines", name=sensor["nome"],
                                 line=dict(color=COR_DESTAQUE, width=2)))

        fora = [(xi, yi) for xi, yi in zip(x, y)
                if (minimo is not None and yi < minimo) or (maximo is not None and yi > maximo)]
        if fora:
            fig.add_trace(go.Scatter(x=[p[0] for p in fora], y=[p[1] for p in fora], mode="markers",
                                     name="Fora do limite", marker=dict(color="#F0616D", size=6)))

        fig.add_trace(go.Scatter(x=[x[-1]], y=[y[-1]], mode="markers", showlegend=False, hoverinfo="skip",
                                 marker=dict(color=COR_DESTAQUE, size=10, line=dict(color=COR_TEXTO, width=1.5))))

        for chave, rotulo in (("min", "Limite mínimo"), ("max", "Limite máximo")):
            valor_limite = limites.get(chave)
            if valor_limite is None:
                continue
            fig.add_trace(go.Scatter(
                x=[x[0], x[-1]], y=[valor_limite, valor_limite], mode="lines",
                name=f"{rotulo} ({valor_limite:g} {sensor['unidade']})".strip(),
                line=dict(color="#F0616D", dash="dash", width=1.5)))

        if sensor.get("atencao_max") is not None:
            valor_atencao = sensor["atencao_max"]
            fig.add_trace(go.Scatter(
                x=[x[0], x[-1]], y=[valor_atencao, valor_atencao], mode="lines",
                name=f"Atenção ({valor_atencao:g} {sensor['unidade']})".strip(),
                line=dict(color="#D4AF37", dash="dot", width=1.5)))

    fig.update_layout(
        title=dict(text=titulo, x=0, xanchor="left", font=dict(size=14, color=COR_SECUNDARIA)),
        height=360, margin=dict(l=48, r=16, t=60 if sensor["id"] == "mq135" else 44, b=60),
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
        font=dict(color=COR_SECUNDARIA, family="Segoe UI"),
        showlegend=True, legend=dict(orientation="h", y=-0.2),
        hovermode="x unified", uirevision=sensor["id"],
    )
    fig.update_xaxes(gridcolor=COR_BORDA, zeroline=False)
    fig.update_yaxes(gridcolor=COR_BORDA, zeroline=False, title=sensor["unidade"])
    if "faixa" in sensor:
        fig.update_yaxes(range=list(sensor["faixa"]))
    elif sensor.get("eixo_zero"):
        fig.update_yaxes(rangemode="tozero")
    if not tem_dados:
        fig.add_annotation(text="Sem dados", x=0.5, y=0.5, xref="paper", yref="paper",
                           showarrow=False, font=dict(size=20, color=COR_SECUNDARIA))
    return fig


def criar_painel_grafico(sensor):
    """Um gráfico aberto: botão Fechar em cima e o gráfico do sensor embaixo.
    Os ids seguem o padrão {"type": ..., "id": sensor} para os callbacks acharem cada um."""
    return html.Section([
        html.Div(html.Button("Fechar ✕", id={"type": "fechar", "id": sensor["id"]}, n_clicks=0,
                             className="botao-mini", title="Fechar este gráfico"),
                 className="g-topo"),
        dcc.Graph(id={"type": "grafico", "id": sensor["id"]}, figure=criar_figura(sensor),
                  config={"displayModeBar": False}),
    ], className="painel g-painel")


def criar_area_graficos(abertos):
    """Conteúdo da área de gráficos: mensagem se não há nenhum aberto, ou um painel por sensor."""
    if not abertos:
        return [html.Div("Clique em um sensor acima para abrir o gráfico dele. Você pode abrir vários ao mesmo tempo.",
                         className="vazio")]
    return [criar_painel_grafico(SENSOR_POR_ID[sensor_id]) for sensor_id in abertos if sensor_id in SENSOR_POR_ID]


# ---------------------------------------------------------------
# PÁGINA
# ---------------------------------------------------------------
app = Dash(__name__, suppress_callback_exceptions=True)
app.title = "MCI | Central de Monitoramento"

NOTA_TECNICA = ("Nota técnica: o ppm do gás é uma ESTIMATIVA pela curva do datasheet do MQ-135 (equivalente de amônia), depende da "
                "calibração do R0 em ar limpo e muda com temperatura e umidade. O sensor não identifica qual gás é. "
                "Limites do gás: NR-15, pior caso entre os gases candidatos. Limites de BPM, SpO₂ e temperatura são provisórios. "
                "BPM e SpO₂ não substituem equipamento médico. Protótipo escolar, não certificado para segurança industrial.")

nomes_simulados = [s["nome"] for s in SENSORES if s["origem"] == "simulado"]
aviso_simulado = html.Div(
    f"MODO DEMONSTRAÇÃO: os dados de {', '.join(nomes_simulados)} são simulados e NÃO são leituras reais.",
    className="faixa-alerta demo",
) if nomes_simulados else None

aviso_teste = html.Div(
    "MODO TESTE: a placa está SIMULADA. Todos os valores, o GPS e os eventos são falsos e NÃO são leituras reais.",
    className="faixa-alerta demo",
) if MODO_TESTE else None


def criar_aba(id_botao, icone, texto, ativa=False):
    return html.Button([html.Span(className=f"icone icone-{icone}"), html.Span(texto)],
                       id=id_botao, n_clicks=0, className="aba ativa" if ativa else "aba")


app.layout = html.Div(
    [
        html.Div([
            html.Header([
                html.Div([
                    html.Div([html.Div("MCI", className="marca-nome"),
                              html.Div("Monitoramento de condições individuais", className="marca-sub")]),
                    html.Div("Central ativa", className="central-status"),
                ], className="topo-linha"),
                html.Nav([
                    criar_aba("nav-painel", "painel", "Painel principal", ativa=True),
                    criar_aba("nav-gps", "gps", "GPS / Localização"),
                    criar_aba("nav-historico", "hist", "Histórico de eventos"),
                ], className="abas"),
            ], className="cabecalho"),

            html.Div([
                html.Aside([html.Div("Conexão", className="rotulo"),
                            html.Div(id="painel-conexao")], className="lateral"),

                html.Main([
                    html.Div([html.H1("Painel principal"),
                              html.Div(datetime.now().strftime("%d/%m/%Y • %H:%M:%S"), id="hora-painel", className="hora")],
                             className="titulo-linha"),
                    aviso_teste,
                    aviso_simulado,
                    html.Div("Verificando limites...", id="painel-alertas", className="faixa-alerta ok"),
                    html.Div([criar_cartao(s) for s in SENSORES if s["id"] != "gps"], className="grade-cartoes"),
                    html.Div("Nenhum sensor enviando dados no momento.", id="aviso-sem-sensores", className="vazio oculto"),
                    html.Div(criar_area_graficos([]), id="graficos-abertos", className="g-abertos"),
                    html.P(NOTA_TECNICA, className="nota"),
                ], id="main-screen", className="conteudo"),

                html.Main(id="pagina-secundaria", children=[], className="conteudo oculto"),
            ], className="corpo"),
        ], className="pagina"),

        dcc.Store(id="sensores-abertos", data=[]),   # lista dos sensores com gráfico aberto (começa vazia)
        dcc.Interval(id="relogio", interval=1000, n_intervals=0),
        dcc.Interval(id="gps-relogio", interval=5000, n_intervals=0),
        dcc.Interval(id="graficos-relogio", interval=3000, n_intervals=0),
    ]
)


# ---------------------------------------------------------------
# CALLBACKS
# ---------------------------------------------------------------
@app.callback(
    Output("main-screen", "className"),
    Output("pagina-secundaria", "children"),
    Output("pagina-secundaria", "className"),
    Output("nav-painel", "className"),
    Output("nav-gps", "className"),
    Output("nav-historico", "className"),
    Input("nav-painel", "n_clicks"),
    Input("nav-gps", "n_clicks"),
    Input("nav-historico", "n_clicks"),
)
def navegar_telas(click_painel, click_gps, click_historico):
    """Mantém GPS e histórico em telas próprias, ocupando toda a área de conteúdo."""
    acionado = ctx.triggered_id
    if acionado == "nav-gps":
        pagina = "gps"
    elif acionado == "nav-historico":
        pagina = "historico"
    elif acionado == "nav-painel":
        pagina = "painel"
    else:
        # A atualização periódica não deve retirar o usuário da tela atual.
        pagina = getattr(navegar_telas, "pagina_atual", "painel")

    navegar_telas.pagina_atual = pagina
    classe_principal = "conteudo" if pagina == "painel" else "conteudo oculto"
    classe_secundaria = "conteudo oculto" if pagina == "painel" else "conteudo"
    abas = ["aba ativa" if pagina == p else "aba" for p in ("painel", "gps", "historico")]

    if pagina == "gps":
        conteudo = criar_tela_gps()
    elif pagina == "historico":
        conteudo = criar_tela_historico()
    else:
        conteudo = []
    return classe_principal, conteudo, classe_secundaria, *abas


@app.callback(
    Output("mapa-gps-separado", "figure"),
    Output("gps-resumo", "children"),
    Input("gps-relogio", "n_intervals"),
    prevent_initial_call=True,
)
def atualizar_mapa_gps(n):
    # Atualiza o mapa e a frase de resumo; a página não é desmontada a cada vez.
    return criar_mapa(), texto_resumo_gps()


@app.callback(
    Output("tabela-historico-aba", "children"),
    Output("historico-total", "children"),
    Output("historico-alertas-total", "children"),
    Output("historico-info-total", "children"),
    Output("historico-grafico-horas", "figure"),
    Output("historico-grafico-dias", "figure"),
    Input("relogio", "n_intervals"),
    prevent_initial_call=True,
)
def atualizar_tela_historico(n):
    ocorrencias = copiar_historico()
    fig_horas, fig_dias = criar_graficos_historico()
    return (criar_tabela_historico(), str(len(ocorrencias)),
            str(sum(1 for o in ocorrencias if o.get("alerta"))),
            str(sum(1 for o in ocorrencias if not o.get("alerta"))), fig_horas, fig_dias)


@app.callback(
    Output("download-historico", "data"),
    Output("aviso-salvar", "children"),
    Input("botao-salvar-excel", "n_clicks"),
    prevent_initial_call=True,
)
def salvar_historico(cliques):
    """Botão 'Salvar como Excel': gera o .xlsx e o navegador baixa o arquivo."""
    if not cliques:
        raise PreventUpdate   # botão recém-criado, não é clique de verdade
    ocorrencias = copiar_historico()
    if not ocorrencias:
        return no_update, "Ainda não há ocorrências para salvar."
    nome = f"historico_mca_{datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.xlsx"
    aviso = f"Arquivo '{nome}' gerado com {len(ocorrencias)} ocorrências (o navegador salva na pasta Downloads)."
    return dcc.send_bytes(historico_para_xlsx(ocorrencias), nome), aviso


@app.callback(
    Output("sensores-abertos", "data"),
    Input({"type": "cartao", "id": ALL}, "n_clicks"),
    Input({"type": "fechar", "id": ALL}, "n_clicks"),
    State("sensores-abertos", "data"),
    prevent_initial_call=True,
)
def alternar_grafico(cliques_cartoes, cliques_fechar, abertos):
    """Clicou num cartão (ou em 'Fechar'): abre o gráfico desse sensor ou, se já estava aberto, fecha."""
    if not ctx.triggered or not ctx.triggered[0]["value"]:
        raise PreventUpdate   # componente recém-criado, não é clique de verdade
    sensor_id = ctx.triggered_id["id"]
    abertos = list(abertos or [])
    if sensor_id in abertos:
        abertos.remove(sensor_id)
    else:
        abertos.append(sensor_id)
    return abertos


@app.callback(
    Output("graficos-abertos", "children"),
    Input("sensores-abertos", "data"),
)
def montar_graficos(abertos):
    """Só monta/remove os painéis quando a lista de abertos muda."""
    return criar_area_graficos(abertos or [])


@app.callback(
    Output({"type": "grafico", "id": ALL}, "figure"),
    Input("graficos-relogio", "n_intervals"),
    State({"type": "grafico", "id": ALL}, "id"),
)
def atualizar_graficos(n, ids):
    """A cada 3 s redesenha os gráficos que estão abertos (sem recriar os painéis)."""
    return [criar_figura(SENSOR_POR_ID[item["id"]]) for item in ids]


@app.callback(
    Output({"type": "valor", "id": ALL}, "children"),
    Output({"type": "estado", "id": ALL}, "children"),
    Output({"type": "extra", "id": ALL}, "children"),
    Output({"type": "cartao", "id": ALL}, "className"),
    Output("aviso-sem-sensores", "className"),
    Input("relogio", "n_intervals"),
    Input("sensores-abertos", "data"),
)
def atualizar_cartoes(n, abertos):
    """A cada segundo atualiza número, estado, extras e classe (cor / visível ou oculto) de cada cartão."""
    abertos = abertos or []
    valores, textos, extras, classes = [], [], [], []
    algum_ativo = False
    for sensor in SENSORES:
        if sensor["id"] == "gps":
            continue  # o GPS tem aba própria
        estado = estado_do_sensor(sensor)
        numero, _, unidade = texto_valor(sensor).rpartition(" ")
        valores.append([html.Span(numero), html.Span(unidade, className="unidade")])
        textos.append(texto_estado(sensor, estado))
        extras.append(criar_extra(sensor))
        classes.append(classe_cartao(estado, sensor["id"] in abertos))
        algum_ativo = algum_ativo or estado not in ESTADOS_INATIVOS
    return valores, textos, extras, classes, "vazio oculto" if algum_ativo else "vazio"


@app.callback(
    Output({"type": "spark", "id": ALL}, "figure"),
    Input("graficos-relogio", "n_intervals"),
)
def atualizar_sparklines(n):
    return [criar_spark(s) for s in SENSORES if s["id"] != "gps"]


@app.callback(
    Output("painel-alertas", "children"),
    Output("painel-alertas", "className"),
    Input("relogio", "n_intervals"),
)
def atualizar_painel_alertas(n):
    """Faixa discreta quando está tudo bem; destacada quando há sensor fora do limite."""
    frases = lista_de_alertas()
    if not frases:
        return "Nenhum alerta ativo", "faixa-alerta ok"
    return [html.Strong("Atenção"), *[html.Span(f) for f in frases]], "faixa-alerta alerta"


@app.callback(
    Output("painel-conexao", "children"),
    Input("relogio", "n_intervals"),
)
def atualizar_conexao(n):
    """Mostra só quem está enviando dados. Sensores com origem None (ainda não ligados)
    não entram na contagem de 'previstos'."""
    previstos = [s for s in SENSORES if s["origem"]]
    ativos = [s for s in previstos if sensor_ativo(s)]
    contador = html.Div([html.Span(str(len(ativos)), className="conexao-numero"),
                         html.Span(f" de {len(previstos)} sensores ativos", className="conexao-texto")])
    if not STATUS_SERIAL["conectado"]:
        return [html.Div(STATUS_SERIAL["mensagem"], className="conexao-aviso"), contador]
    return [html.Div(STATUS_SERIAL["mensagem"], className="conexao-info"), contador,
            html.Ul([html.Li(s["nome"]) for s in ativos], className="conexao-lista")]


@app.callback(
    Output("hora-painel", "children"),
    Input("relogio", "n_intervals"),
)
def atualizar_hora_painel(n):
    return datetime.now().strftime("%d/%m/%Y • %H:%M:%S")


if __name__ == "__main__":
    if CALIBRAR:
        calibrar_r0()
        sys.exit(0)
    iniciar_threads()
    if MODO_TESTE:
        print("[central] MODO TESTE: placa simulada. Os dados são FALSOS.")
    elif MQ135_R0_KOHM is None:
        print("[central] Gás sem calibração: rode  python central_mca.py --calibrar  e copie o R0 para MQ135_R0_KOHM.")
    print("[central] Abra no navegador: http://127.0.0.1:8050")
    # use_reloader=False: evita rodar o arquivo duas vezes (duas threads na mesma COM).
    app.run(debug=True, use_reloader=False)